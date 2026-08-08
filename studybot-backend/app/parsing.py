import io

import pymupdf  
from docx import Document as DocxDocument
from pptx import Presentation


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_parts = []
    with pymupdf.open(stream=file_bytes, filetype="pdf") as pdf:
        for page in pdf:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs]
    return "\n".join(paragraphs)


def extract_text_from_pptx(file_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(file_bytes))
    text_parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line:
                        text_parts.append(line)
    return "\n".join(text_parts)


def extract_text_from_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore")


# --- Registry: maps MIME type -> extraction function ---
# To add legacy .doc/.ppt support later: write extract_text_from_legacy_doc()
# etc., add the corresponding MIME type + function pair below, and add the
# MIME type back to ALLOWED_MIME_TYPES in .env. No other code changes needed.
PARSERS = {
    "application/pdf": extract_text_from_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": extract_text_from_docx,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": extract_text_from_pptx,
    "text/plain": extract_text_from_txt,
}


def extract_text(file_bytes: bytes, mime_type: str) -> str:
    """
    Dispatches to the correct parser based on MIME type.
    Raises ValueError if the MIME type has no registered parser —
    this should never actually happen in practice since the upload
    endpoint already validates against ALLOWED_MIME_TYPES, but we
    fail loudly here rather than silently returning empty text.
    """
    parser = PARSERS.get(mime_type)
    if parser is None:
        raise ValueError(f"No parser registered for MIME type: {mime_type}")

    text = parser(file_bytes)
    return text.strip()